"""
test_backend_full.py
Comprehensive test suite for all app backend services and algorithms.
"""

import os
import sys
import tempfile
import pathlib

# Setup test DB path in tempdir before importing config
test_db_dir = tempfile.mkdtemp()
test_db_path = os.path.join(test_db_dir, "test_tax.db")
os.environ["DATABASE_PATH"] = test_db_path
os.environ["DRY_RUN_MODE"] = "true"

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

from database.db import init_db, get_connection
from services import (
    resident_service,
    settings_service,
    template_service,
    word_import_service,
    excel_service,
    reminder_engine,
    messaging_service,
)
import main

def test_all():
    print("1. Initializing DB...")
    init_db()
    
    # Verify seeded cycle
    cycle = settings_service.get_active_cycle()
    assert cycle is not None, "Active cycle should be seeded"
    print(f"   Active cycle: FY {cycle['fy_label']}, due: {cycle['due_date']}")

    # Verify seeded cadence
    cadence = settings_service.get_cadence(cycle['id'])
    assert len(cadence) > 0, "Cadence should be seeded"
    print(f"   Seeded cadence rows: {len(cadence)}")

    # Verify seeded templates
    templates = template_service.get_all_templates()
    assert len(templates) >= 8, f"Expected at least 8 templates, got {len(templates)}"
    print(f"   Seeded templates: {len(templates)}")

    # 2. Test Resident Service CRUD
    print("\n2. Testing Resident Service CRUD...")
    r1 = resident_service.create_resident({
        "name": "Anil Deshmukh",
        "property_id": "GP/2026/001",
        "ward": "Ward 1",
        "phone": "9876543210",
        "address": "House 12",
        "base_amount": 2000.0,
        "payment_status": "unpaid",
    })
    assert r1["id"] is not None
    print(f"   Created resident id={r1['id']} name={r1['name']}")

    r2 = resident_service.create_resident({
        "name": "Sunita Patil",
        "property_id": "GP/2026/002",
        "ward": "Ward 2",
        "phone": "9123456789",
        "address": "Plot 5",
        "base_amount": 3500.0,
        "payment_status": "unpaid",
    })

    # Duplicate property_id check
    try:
        resident_service.create_resident({
            "name": "Duplicate Person",
            "property_id": "GP/2026/001",
            "phone": "9000000000",
            "base_amount": 1000.0,
        })
        assert False, "Should have raised ValueError for duplicate property_id"
    except ValueError:
        print("   Duplicate property_id correctly rejected")

    # Update resident
    updated = resident_service.update_resident(r1["id"], {
        "name": "Anil R. Deshmukh",
        "ward": "Ward 1A",
    })
    assert updated["name"] == "Anil R. Deshmukh"
    assert updated["ward"] == "Ward 1A"
    print("   Resident updated successfully")

    # Mark paid
    paid_r1 = resident_service.mark_paid(r1["id"], "2026-01-15")
    assert paid_r1["payment_status"] == "paid"
    assert paid_r1["paid_date"] == "2026-01-15"
    print("   Mark paid works")

    # Mark unpaid
    unpaid_r1 = resident_service.mark_unpaid(r1["id"])
    assert unpaid_r1["payment_status"] == "unpaid"
    assert unpaid_r1["paid_date"] is None
    print("   Mark unpaid works")

    # Stats check
    stats = resident_service.get_stats(cycle)
    assert stats["total"] == 2
    assert stats["unpaid"] == 2
    print(f"   Stats check: total={stats['total']}, pending_base={stats['pending_base']}")

    # 3. Penalty and Rebate calculation tests
    print("\n3. Testing Penalty and Rebate calculations...")
    cycle_flat = {
        "penalty_type": "flat",
        "penalty_value": 50.0,
        "penalty_start_days": 1,
        "rebate_enabled": 1,
        "rebate_percent": 5.0,
        "rebate_deadline": "2099-12-31",
    }
    pen_flat_0 = settings_service.calculate_penalty(2000, cycle_flat, 0)
    assert pen_flat_0 == 0.0
    pen_flat_15 = settings_service.calculate_penalty(2000, cycle_flat, 15)
    assert pen_flat_15 == 50.0  # 1 month late
    pen_flat_45 = settings_service.calculate_penalty(2000, cycle_flat, 45)
    assert pen_flat_45 == 100.0  # 2 months late
    reb_5 = settings_service.calculate_rebate(2000, cycle_flat)
    assert reb_5 == 100.0  # 5% of 2000
    print("   Flat penalty & rebate calculations verified")

    cycle_pct = {
        "penalty_type": "percent",
        "penalty_value": 2.0,  # 2% per month
        "penalty_start_days": 1,
        "rebate_enabled": 0,
    }
    pen_pct_45 = settings_service.calculate_penalty(2000, cycle_pct, 45)
    assert pen_pct_45 == 80.0  # 2000 * 2% * 2 months = 80.0
    print("   Percentage penalty calculation verified")

    # 4. Word Import tests
    print("\n4. Testing Word Import integration...")
    raw_paste = """
Name: Ganesh Shinde
Property ID: GP/2026/003
Ward: Ward 3
Phone: 9988776655
Amount: ₹4,000

Name: Anil R. Deshmukh
Property ID: GP/2026/001
Ward: Ward 1
Phone: 9876543210
Amount: 2200
"""
    parse_res = main.parse_import_text(raw_paste)
    assert parse_res["success"] is True
    assert len(parse_res["new"]) == 1
    assert parse_res["new"][0]["property_id"] == "GP/2026/003"
    assert len(parse_res["duplicates"]) == 1
    assert parse_res["duplicates"][0]["property_id"] == "GP/2026/001"
    print(f"   parse_import_text: {len(parse_res['new'])} new, {len(parse_res['duplicates'])} duplicates")

    conf_res = main.confirm_import(parse_res["new"])
    assert conf_res["success"] is True
    assert conf_res["inserted"] == 1
    print("   confirm_import inserted new record")

    # Real .docx base64 upload test
    import docx, base64, io
    test_doc = docx.Document()
    test_doc.add_paragraph("Name: Vijay Docx\nProperty ID: GP/2026/099\nWard: Ward 4\nPhone: 9876543211\nAmount: 3500")
    test_bio = io.BytesIO()
    test_doc.save(test_bio)
    docx_b64 = "data:application/vnd.openxmlformats-officedocument.wordprocessingml.document;base64," + base64.b64encode(test_bio.getvalue()).decode("utf-8")
    docx_res = main.parse_import_docx(docx_b64)
    assert docx_res["success"] is True
    assert len(docx_res["new"]) == 1
    assert docx_res["new"][0]["property_id"] == "GP/2026/099"
    print("   parse_import_docx base64 file upload verified")

    merge_res = main.merge_resident_import(r1["id"], {"name": "Anil Deshmukh (Merged)", "ward": "Ward 1", "phone": "9876543210", "address": "", "base_amount": 2200.0})
    assert merge_res["merged"] is True
    r1_merged = resident_service.get_resident_by_id(r1["id"])
    assert r1_merged["name"] == "Anil Deshmukh (Merged)"
    assert r1_merged["base_amount"] == 2200.0
    print("   merge_resident_import merged record")

    # 4b. Excel Import and Template tests
    print("\n4b. Testing Excel Import & Template integration...")
    tpl_res = main.get_excel_import_template()
    assert tpl_res["success"] is True
    assert "Taxpayer_Import_Template.xlsx" in tpl_res["filename"]
    print("   get_excel_import_template verified")

    # Generate real Excel file and parse with main.parse_import_excel
    import openpyxl
    wb_test = openpyxl.Workbook()
    ws_test = wb_test.active
    ws_test.append(["Taxpayer Name", "Property ID", "Ward", "Phone", "Base Tax Amount", "Address"])
    ws_test.append(["Excel Test User", "GP/2026/777", "Ward 5", "9900112233", 4200, "Plot 99, North Street"])
    ws_test.append(["Anil Deshmukh (Merged)", "GP/2026/001", "Ward 1", "9876543210", 2200, "House 12"])
    xlsx_bio = io.BytesIO()
    wb_test.save(xlsx_bio)
    xlsx_b64 = "data:application/vnd.openxmlformats-officedocument.spreadsheetml.sheet;base64," + base64.b64encode(xlsx_bio.getvalue()).decode("utf-8")

    excel_res = main.parse_import_excel(xlsx_b64)
    assert excel_res["success"] is True
    assert len(excel_res["new"]) == 1
    assert excel_res["new"][0]["property_id"] == "GP/2026/777"
    assert len(excel_res["duplicates"]) == 1
    assert excel_res["duplicates"][0]["property_id"] == "GP/2026/001"
    print("   parse_import_excel (.xlsx file) verified")

    # 5. Reminder Engine and Dedup tests
    print("\n5. Testing Reminder Engine & Dedup...")
    # Preview message
    prev_mr = reminder_engine.preview_message(r1["id"], "upcoming", "mr")
    assert "Anil Deshmukh (Merged)" in prev_mr
    assert "2,200" in prev_mr
    prev_en = reminder_engine.preview_message(r1["id"], "upcoming", "en")
    assert "Anil Deshmukh (Merged)" in prev_en
    print(f"   preview_message generated:\n     {prev_en.splitlines()[0]}")

    # Send reminder now (since date is past due date, stage is post_due)
    send_res = reminder_engine.send_now(r1["id"], "post_due", force=False)
    assert send_res.get("status") == "sent"
    print("   send_now succeeded in dry-run mode")

    # Dedup check - check already_sent returns True for the sent offset
    assert reminder_engine.already_sent(r1["id"], cycle["id"], "post_due", 3) is True
    print("   Dedup check verified: already_sent is True for offset 3")

    # Send with force=True should proceed
    send_res_force = reminder_engine.send_now(r1["id"], "post_due", force=True)
    assert send_res_force.get("status") == "sent"
    print("   Force send verified")

    # Check reminder log
    log_rows = reminder_engine.get_reminder_log(r1["id"])
    assert len(log_rows) >= 1
    print(f"   Reminder log verified: {len(log_rows)} rows for resident")

    # Test run_daily_check
    daily_res = reminder_engine.run_daily_check()
    assert daily_res["status"] == "done"
    print(f"   Daily check run summary: {daily_res}")

    # 6. Report data tests
    print("\n6. Testing Report Data & Excel Export queries...")
    rep_all = main.get_report_data("all", "", "all")
    assert rep_all["success"] is True
    assert len(rep_all["records"]) == 3
    print(f"   Report all: {len(rep_all['records'])} records, total outstanding={rep_all['total_outstanding']}")

    rep_excel_res = main.generate_report_excel("all", "", "all")
    assert rep_excel_res["success"] is True
    assert "Tax_Report_all_All.xlsx" in rep_excel_res["filename"]
    assert "base64," in rep_excel_res["data"]
    print(f"   generate_report_excel verified: {rep_excel_res['filename']}")

    print("\n=== ALL TESTS PASSED SUCCESSFULLY! ===")

if __name__ == "__main__":
    test_all()
